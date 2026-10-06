"""Build the designed handbook from Markdown and genuine mobile screenshots.
Dependencies: python-docx, Pillow. Optional PDF QA: PyMuPDF.
Install locally: python -m pip install --target .handbook-tools python-docx Pillow PyMuPDF
"""
from pathlib import Path
import argparse
import sys
import re
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '.handbook-tools'))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--school', action='store_true', help='Build the school handbook')
parser.add_argument('--output', type=Path, help='Alternative DOCX path, relative to the repository or absolute')
args = parser.parse_args()
import subprocess
subprocess.run(['node', str(ROOT / 'scripts/handbook-screenshots.mjs')], cwd=ROOT, check=True)
SCHOOL = args.school
VERSION = '1.5' if SCHOOL else '1.6'
STEM = 'Betriebshandbuch-Flugschulen' if SCHOOL else 'Benutzerhandbuch-Piloten'
SOURCE = ROOT / f'docs/{STEM}.md'
DEST = (ROOT / args.output).resolve() if args.output else ROOT / f'docs/{STEM}-v{VERSION}.docx'
NAVY, TEAL, GRAY, PALE = '0F2338', '0A6E99', '576A7A', 'F0F5F9'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2.0), Cm(1.8)
section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)
section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
section.different_first_page_header_footer = True
styles = doc.styles
for name in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    st = styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(NAVY)
    st.paragraph_format.space_after = Pt(5 if SCHOOL else 6)
    st.paragraph_format.line_spacing = 1.05 if SCHOOL else 1.12
    st.paragraph_format.widow_control = True
    if SCHOOL: st.paragraph_format.keep_together = True
styles['Heading 1'].font.name = 'Segoe UI'
styles['Heading 1'].font.size = Pt(23)
styles['Heading 1'].font.color.rgb = RGBColor.from_string(NAVY)
styles['Heading 1'].paragraph_format.space_after = Pt(15)
styles['Heading 1'].paragraph_format.keep_with_next = True
styles['Heading 2'].font.name = 'Calibri'
styles['Heading 2'].font.size = Pt(13)
styles['Heading 2'].font.color.rgb = RGBColor.from_string(TEAL)
styles['Heading 2'].paragraph_format.space_before = Pt(12)
styles['Heading 2'].paragraph_format.space_after = Pt(6)
styles['Heading 2'].paragraph_format.keep_with_next = True
styles['Caption'].font.name = 'Calibri'
styles['Caption'].font.size = Pt(9)
styles['Caption'].font.color.rgb = RGBColor.from_string(GRAY)
styles['Caption'].paragraph_format.line_spacing = 1.1

def shade(cell, fill):
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(sh)

def rule(paragraph, color=TEAL, size='10'):
    borders = OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
    for k,v in {'val':'single','sz':size,'space':'7','color':color}.items(): bottom.set(qn('w:'+k),v)
    borders.append(bottom); paragraph._p.get_or_add_pPr().append(borders)

def inline(p, text):
    for part in re.split(r'(\*\*.*?\*\*)', text):
        run=p.add_run(part[2:-2] if part.startswith('**') else part)
        run.bold=part.startswith('**')
    return p

def label(text, target=doc):
    p=target.add_paragraph()
    p.paragraph_format.space_after=Pt(8)
    r=p.add_run(text.upper()); r.font.size=Pt(9); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
    return p

def field(p, instruction):
    r=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); r._r.append(begin)
    txt=OxmlElement('w:instrText'); txt.set(qn('xml:space'),'preserve'); txt.text=instruction; r._r.append(txt)
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); r._r.append(sep)
    t=OxmlElement('w:t'); t.text='1'; r._r.append(t)
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.append(end)

header=section.header.paragraphs[0]
header.text='FLYARY  /  ' + ('BETRIEBSHANDBUCH FLUGSCHULEN' if SCHOOL else 'PILOTENHANDBUCH')
header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
rule(header,'D7E4EB','4')
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run(f'Version {VERSION} · 6. Oktober 2026   ·   ')
field(footer,' PAGE ')
for r in footer.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

# Editorial cover: clear hierarchy, restrained brand colour, no stock imagery.
label('Flyary · Anwenderdokumentation')
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(38)
r=p.add_run('Flyary'); r.font.name='Segoe UI'; r.font.size=Pt(52); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(16)
r=p.add_run('Deine Flugschule.\nDein Überblick.' if SCHOOL else 'Dein Flugbuch.\nDein Fortschritt.'); r.font.size=Pt(29); r.font.bold=True
rule(p)
p=doc.add_paragraph('Betriebshandbuch für\nFlugschulen' if SCHOOL else 'Benutzerhandbuch für\nPilotinnen und Piloten'); p.paragraph_format.space_before=Pt(20)
for r in p.runs: r.font.size=Pt(18)
table=doc.add_table(rows=1, cols=2); table.autofit=False; table.columns[0].width=Cm(9.8); table.columns[1].width=Cm(7.2)
left,right=table.rows[0].cells
left.width=Cm(9.8); right.width=Cm(7.2)
left.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
left.paragraphs[0].text='Von der Tagesplanung bis zur\nAusbildung und Abrechnung.' if SCHOOL else 'Vom ersten Einstieg bis zum\nvollständigen Flugbucheintrag.'
for r in left.paragraphs[0].runs: r.font.size=Pt(15)
left.add_paragraph('Schritt für Schritt erklärt.\nMit mobilen Bildschirmansichten,\nPraxisbeispielen und Fehlerhilfe.')
label(f'Version {VERSION} · 6. Oktober 2026',left)
right.paragraphs[0].add_run().add_picture(str(ROOT/('docs/handbook/school-mobile/01-overview.png' if SCHOOL else 'docs/handbook/mobile/01-home.png')),width=Cm(5.3))

text=SOURCE.read_text(encoding='utf-8')
parts=re.split(r'^## (.+)$',text,flags=re.M)
label('Über dieses Handbuch').paragraph_format.page_break_before=True
doc.add_paragraph('Sicher durch die App',style='Title')
for paragraph in parts[0].split('\n\n')[3:]:
    if paragraph.strip(): inline(doc.add_paragraph(),paragraph.strip())
doc.add_paragraph('Bild und Anleitung gehören zusammen',style='Heading 2')
doc.add_paragraph('Die Bildschirmansichten stehen unmittelbar bei den zugehörigen Arbeitsschritten. Links siehst du die mobile App, rechts die Reihenfolge der Aktionen und die erwarteten Eingaben. Pflichtangaben und optionale Ergänzungen werden im Ablauf benannt. Gezeigte Namen, Termine und Flugwerte sind fiktiv; die Oberfläche stammt aus der echten App.')
doc.add_paragraph('Was dieses Handbuch abdeckt',style='Heading 2')
doc.add_paragraph('Beschrieben ist die Bedienung von Flyary durch Schulleitung, Fluglehrer, Starthelfer und Administration. Funktionen hängen von Schulmitgliedschaft und Berechtigung ab. Dieses Handbuch erklärt die Softwareabläufe; schuleigene Ausbildungs-, Sicherheits- und Betriebsverfahren werden darin nicht festgelegt.' if SCHOOL else 'Beschrieben ist die Pilotensicht. Funktionen hängen von Datenbestand, Gruppenmitgliedschaft und Berechtigung ab. Die Ausbildung und Administration einer Flugschule werden hier nicht vollständig dokumentiert.')
label('Orientierung').paragraph_format.page_break_before=True
doc.add_paragraph('Inhalt',style='Title')
p=doc.add_paragraph()
field(p,' TOC \\o "1-1" \\h \\z \\u ')
doc.add_paragraph('In Word sind die Kapitel im Inhaltsverzeichnis und über den Navigationsbereich erreichbar. Die Seitenzahlen werden bei der Dokumenterstellung aktualisiert.',style='Caption')

figure_no=0
def walkthrough(title, lines, supporting=None):
    global figure_no
    figure_no += 1
    picture = next(re.match(r'!\[(.*?)\]\((.*?)\)', line) for line in lines if line.startswith('!['))
    note = ' '.join(line[1:].strip() for line in lines if line.startswith('>'))
    steps = [line for line in lines if re.match(r'^\d+\. ',line)]
    assert steps, title
    p=doc.add_paragraph(title,style='Heading 2')
    p.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=3);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(table.columns,[7.4,.4,9.2]):c.width=Cm(w)
    for c,w in zip(table.rows[0].cells,[7.4,.4,9.2]):c.width=Cm(w)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    left,right=table.cell(0,0),table.cell(0,2)
    left.vertical_alignment=right.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    img=ROOT/'docs'/picture[2]
    with Image.open(img) as im:assert im.size==(780,1688),str(img)
    pic=left.paragraphs[0].add_run().add_picture(str(img),width=Cm(7.0))
    pic._inline.docPr.set('descr',picture[1]+' – Smartphone-Ansicht, fiktive Beispieldaten')
    cp=left.add_paragraph(f'Abb. {figure_no:02d} · {picture[1]}',style='Caption')
    cp.paragraph_format.space_before=Pt(5)
    if note:left.add_paragraph(note,style='Caption')
    rp=right.paragraphs[0];rp.text='SCHRITT FÜR SCHRITT'
    for run in rp.runs:run.font.size=Pt(9);run.font.bold=True;run.font.color.rgb=RGBColor.from_string(TEAL)
    for step in steps:
        p=inline(right.add_paragraph(),step)
        p.paragraph_format.space_after=Pt(10)
        p.paragraph_format.left_indent=Cm(.45);p.paragraph_format.first_line_indent=Cm(-.45)
    # Keep short explanatory notes with the final procedure, using the otherwise
    # empty space beside the phone instead of creating orphan continuation pages.
    if supporting:
        for line in supporting:
            if not line.strip(): continue
            heading=line.startswith('### ')
            p=inline(right.add_paragraph(),line[4:] if heading else line)
            p.paragraph_format.space_after=Pt(5)
            p.paragraph_format.line_spacing=1.03
            if heading:p.paragraph_format.keep_with_next=True
            for run in p.runs:
                run.font.size=Pt(10 if heading else 9.5)
                if heading:
                    run.bold=True;run.font.color.rgb=RGBColor.from_string(TEAL)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

for idx in range(1,len(parts),2):
    title,body=parts[idx],parts[idx+1]
    doc.add_paragraph(title,style='Heading 1').paragraph_format.page_break_before=True
    lines=body.strip().splitlines(); i=0; imgs=[]
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('::: schritte '):
            block_title=line[len('::: schritte '):];block=[];i+=1
            while i<len(lines) and lines[i].strip()!=':::':
                block.append(lines[i].strip());i+=1
            supporting=[]
            tail=lines[i+1:]
            if SCHOOL and not any(s.startswith(':::') or s.startswith('|') for s in tail) and len(' '.join(tail).split())<=150:
                supporting=tail
                i=len(lines)
            walkthrough(block_title,block,supporting)
        elif line.startswith('### '): doc.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line); cap=[]
            while i+1<len(lines) and (not lines[i+1].strip() or lines[i+1].startswith('>')):
                i+=1
                if lines[i].startswith('>'):cap.append(lines[i][1:].strip())
            imgs.append((m[1],m[2],' '.join(cap)))
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',v) for v in cells):rows.append(cells)
                i+=1
            i-=1
            t=doc.add_table(rows=0,cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER
            for ri,row in enumerate(rows):
                cs=t.add_row().cells
                for cell,value in zip(cs,row):
                    p=inline(cell.paragraphs[0],value); p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    for r in p.runs:r.font.size=Pt(9.5)
                    p.paragraph_format.keep_with_next=ri<len(rows)-1
                    if ri==0:
                        shade(cell,NAVY)
                        for r in p.runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
                    elif ri%2:shade(cell,PALE)
                trpr=t.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
                if ri==0:trpr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif re.match(r'^\d+\. ',line):
            p=inline(doc.add_paragraph(),line)
            p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5)
            if SCHOOL:
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=1.0
        else:
            p=inline(doc.add_paragraph(),line)
            if line.startswith('**') and not line.startswith('**Weg:'):
                p.paragraph_format.space_before=Pt(5)
        i+=1
    assert not imgs, 'All images must belong to a walkthrough'

settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'de-CH')
styles['Normal'].element.get_or_add_rPr().append(lang)
doc.core_properties.title='Flyary – Betriebshandbuch für Flugschulen' if SCHOOL else 'Flyary – Benutzerhandbuch für Pilotinnen und Piloten'
doc.core_properties.subject=f'Version {VERSION} – Bebilderte Schritt-für-Schritt-Anleitung'
doc.core_properties.author='Flyary'
doc.core_properties.keywords='Flyary, Flugschule, Betriebshandbuch' if SCHOOL else 'Flyary, Flugbuch, Pilot, Benutzerhandbuch'
doc.save(DEST)
print(f'Created {DEST.name}: {len(text.split())} words, {figure_no} mobile figures, {(len(parts)-1)//2} chapters, version {VERSION}')

